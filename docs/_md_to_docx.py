# -*- coding: utf-8 -*-
"""Convert CHAPITRE_3 markdown to Word (.docx) with embedded diagram PNGs."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(r"c:\Users\hp\Documents\BK CHERIE\hopital-amen\docs")
SRC = BASE / "CHAPITRE_3_ANALYSE_CONCEPTION.md"
OUT = BASE / "CHAPITRE_3_ANALYSE_CONCEPTION.docx"
IMG_WIDTH = Cm(15.5)


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_para(doc, text, size=12, bold=False, italic=False, center=False, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # inline code / bold segments
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, size=size - 1 if size > 10 else size, italic=True)
            run.font.name = "Consolas"
        else:
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level):
    sizes = {1: 16, 2: 14, 3: 13, 4: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 12), bold=True)
    return p


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            run = p.add_run(val)
            set_run_font(run, size=10, bold=(i == 0))
    doc.add_paragraph()


def add_image(doc, img_path: Path, caption: str):
    if not img_path.exists():
        add_para(doc, f"[Image manquante : {img_path.name}]", size=10, italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(img_path), width=IMG_WIDTH)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run_font(r, size=10, italic=True)


def parse_table_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_sep_row(cells):
    return all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    i = 0
    in_code = False
    code_lang = ""
    code_buf = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = [parse_table_row(r) for r in table_buf]
        if len(rows) >= 2 and is_sep_row(rows[1]):
            rows = [rows[0]] + rows[2:]
        add_table(doc, rows)
        table_buf = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_lang = line[3:].strip()
                code_buf = []
            else:
                # skip leftover mermaid/code — images preferred
                label = "Code" if code_lang and code_lang != "mermaid" else None
                if label and code_buf:
                    add_para(doc, f"[{label}]", size=10, italic=True, space_after=4)
                    add_para(doc, "\n".join(code_buf), size=9, space_after=8)
                in_code = False
                code_lang = ""
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Markdown image: ![caption](path)
        m_img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if m_img:
            flush_table()
            caption = m_img.group(1).strip() or "Figure"
            rel = m_img.group(2).strip()
            add_image(doc, (BASE / rel).resolve(), caption)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        else:
            flush_table()

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith("#### "):
            add_heading(doc, line[5:].strip(), 4)
        elif line.startswith("> "):
            add_para(doc, line[2:].strip(), size=11, italic=True)
        elif re.match(r"^[-*] ", line):
            add_para(doc, "• " + line[2:].strip(), size=12, space_after=4)
        elif re.match(r"^\d+\. ", line):
            add_para(doc, line.strip(), size=12, space_after=4)
        else:
            # strip residual markdown links [text](url) -> text (url)
            cleaned = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", line.strip())
            if cleaned.startswith("*") and cleaned.endswith("*") and not cleaned.startswith("**"):
                add_para(doc, cleaned.strip("*"), size=11, italic=True)
            else:
                add_para(doc, cleaned, size=12)

        i += 1

    flush_table()
    doc.save(OUT)
    print(f"OK -> {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
